"""Agente do SAP Ariba: lê o evento de compra exportado (Excel, PDF ou Word), monta o preço com impostos e prepara a resposta.

Sem senha e sem API da SAP: o fornecedor baixa o evento no Ariba (RFQ/RFP, "Exportar para Excel" ou PDF) e envia o arquivo.
Passos:
  1. Leitor: identifica o evento (Doc...), comprador, prazo, moeda e os itens (descrição, quantidade, unidade).
  2. Classificador: separa mercadoria e serviço e marca serviços com retenção (mão de obra, profissionais).
  3. Precificador: com o custo de cada item e a margem desejada, calcula o preço com os tributos do regime (por dentro).
  4. Retenções: aplica as retenções entre empresas (CSRF 4,65%, IRRF 1,5% ou 1%, INSS 11%) e mostra o líquido a receber.
  5. Resposta: devolve a planilha preenchida para importar no Ariba e a lista do que conferir antes de enviar.
"""
import io
import re
import unicodedata

SERV = ["servic", "consultor", "manutenc", "instalac", "limpeza", "assessor", "treinament", "mao de obra", "suporte", "projeto", "auditor", "obra"]
MAO = ["mao de obra", "limpeza", "vigilancia", "manutencao predial", "obra", "construc", "reforma"]
PROF = ["consultor", "assessor", "auditor", "contabil", "advocac", "engenharia", "projeto", "treinament"]
REG = {  # tributos por dentro sobre o preço (%), por regime e natureza
    "presumido": {"servico": {"PIS": 0.65, "COFINS": 3.0, "ISS": 5.0}, "mercadoria": {"PIS": 0.65, "COFINS": 3.0, "ICMS": 18.0}},
    "real": {"servico": {"PIS (líquido estimado)": 0.83, "COFINS (líquido estimado)": 3.8, "ISS": 5.0}, "mercadoria": {"PIS (líquido estimado)": 0.83, "COFINS (líquido estimado)": 3.8, "ICMS": 18.0}},
    "simples": {"servico": {"Simples Nacional (efetiva)": 8.0}, "mercadoria": {"Simples Nacional (efetiva)": 6.0}},
}


def norm(s):
    return unicodedata.normalize("NFD", str(s or "").lower()).encode("ascii", "ignore").decode()


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def ler(nome, dados):
    """Devolve (cabeçalho, itens, texto) a partir do arquivo exportado do Ariba."""
    ext = nome.lower().rsplit(".", 1)[-1]
    linhas, itens = [], []
    if ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(dados), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
            linhas += [" ".join(str(c) for c in r if c is not None) for r in rows]
            for i, r in enumerate(rows):  # cabeçalho da tabela de itens
                h = [norm(c) for c in r]
                if any(x in ("name", "nome", "descricao", "description", "item") for x in h) and any(x.startswith(("quant", "qty")) for x in h):
                    col = lambda *ks: next((j for j, x in enumerate(h) if any(x.startswith(k) for k in ks)), None)
                    cn, cq, cu = col("descricao", "description") if col("descricao", "description") is not None else col("name", "nome", "item"), col("quant", "qty"), col("unit", "unidade", "uom")
                    cc = col("custo", "cost", "preco", "price")
                    for rr in rows[i + 1:]:
                        if not rr or cn is None or cn >= len(rr) or not rr[cn]:
                            continue
                        q = _num(rr[cq]) if cq is not None and cq < len(rr) else None
                        if q is None:
                            continue
                        itens.append({"descricao": str(rr[cn]).strip(), "quantidade": q, "unidade": str(rr[cu]).strip() if cu is not None and cu < len(rr) and rr[cu] else "un",
                                      "custo_unit": _num(rr[cc]) if cc is not None and cc < len(rr) else None})
                    break
    elif ext == "pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(dados)) as pdf:
            linhas = "\n".join(p.extract_text() or "" for p in pdf.pages).splitlines()
    elif ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(dados))
        linhas = [p.text for p in d.paragraphs] + [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
    else:
        linhas = dados.decode("utf-8", "ignore").splitlines()
    texto = "\n".join(linhas)
    if not itens:  # PDF/Word: "1 | Descrição | 10 | UN" ou "Descrição ... Qtd: 10 UN"
        for ln in linhas:
            m = re.match(r"^\s*(\d{1,3})[\s.|)-]+(.+?)\s*[|;\s]\s*(\d+(?:[.,]\d+)?)\s*[|;\s]\s*([A-Za-z]{1,6})\b", ln)
            if m and len(m.group(2)) > 3:
                itens.append({"descricao": m.group(2).strip(" |"), "quantidade": _num(m.group(3)), "unidade": m.group(4), "custo_unit": None})
    g = lambda rx: (re.search(rx, texto, re.I) or [None, None])[1]
    cab = {"evento": g(r"\b(Doc\d{6,})\b"), "titulo": g(r"(?:T[íi]tulo|Title|Event Name|Nome do evento)\s*[:\-]\s*(.+)"),
           "comprador": g(r"(?:Comprador|Buyer|Empresa compradora|Customer)\s*[:\-]\s*(.+)"),
           "prazo": g(r"(?:Prazo|Due Date|Data de t[ée]rmino|Encerramento|End Time)\s*[:\-]\s*([0-9/.\-: ]{8,19}(?:\s*[AP]M)?)"),
           "moeda": g(r"\b(BRL|USD|EUR)\b") or "BRL"}
    return cab, itens, texto


def preparar(nome, dados, regime="presumido", margem=15.0, custos=None):
    passos = []
    cab, itens, texto = ler(nome, dados)
    passos.append(f"Leitor: evento {cab['evento'] or 'sem número'}, {len(itens)} item(ns), prazo {cab['prazo'] or 'não encontrado'}")
    linhas, total, ret_tot = [], 0.0, {}
    for i, it in enumerate(itens):
        d = norm(it["descricao"])
        nat = "servico" if any(k in d for k in SERV) else "mercadoria"
        trib = REG.get(regime, REG["presumido"])[nat]
        carga = sum(trib.values())
        custo = (custos or {}).get(i, it.get("custo_unit"))
        preco_unit = round(custo * (1 + margem / 100) / (1 - carga / 100), 2) if custo else None
        valor = (preco_unit or 0) * it["quantidade"]
        total += valor
        ret = {}
        if nat == "servico" and regime != "simples":
            ret["CSRF (PIS, Cofins e CSLL) 4,65%"] = 4.65  # Lei 10.833/2003, art. 30
            ret["IRRF 1,5%" if any(k in d for k in PROF) else "IRRF 1%"] = 1.5 if any(k in d for k in PROF) else 1.0  # RIR/2018, arts. 714 e 716
        if nat == "servico" and any(k in d for k in MAO):
            ret["INSS 11%"] = 11.0  # Lei 8.212/1991, art. 31
        for k, p in ret.items():
            ret_tot[k] = ret_tot.get(k, 0) + valor * p / 100
        linhas.append({**it, "natureza": nat, "tributos": trib, "carga": carga, "custo_unit": custo, "preco_unit": preco_unit, "valor": round(valor, 2), "retencoes": ret})
    passos.append(f"Classificador: {sum(l['natureza'] == 'servico' for l in linhas)} serviço(s) e {sum(l['natureza'] == 'mercadoria' for l in linhas)} mercadoria(s)")
    passos.append(f"Precificador: regime {regime}, margem {margem:g}%, tributos por dentro do preço")
    rt = sum(ret_tot.values())
    passos.append(f"Retenções: R$ {rt:,.2f} retidos pelo comprador".replace(",", "X").replace(".", ",").replace("X", "."))
    conferir = ["Preço unitário de cada item preenchido no Ariba, na mesma moeda do evento",
                "Prazo de entrega e validade da proposta pedidos no evento",
                "Questionário do evento respondido (certidões, compliance, dados bancários)",
                "Condição de pagamento e se o comprador retém CSRF, IRRF e INSS",
                "Envie antes do prazo: o Ariba bloqueia respostas depois do encerramento"]
    if not linhas:
        conferir.insert(0, "Nenhum item foi lido: exporte o evento pelo Ariba em Excel (Download Content) e envie de novo")
    if any(l["preco_unit"] is None for l in linhas):
        conferir.insert(0, "Informe o custo de cada item para o agente calcular o preço")
    return {"cabecalho": cab, "itens": linhas, "total": round(total, 2), "retencoes": {k: round(v, 2) for k, v in ret_tot.items()},
            "liquido": round(total - rt, 2), "conferir": conferir, "passos": passos, "regime": regime, "margem": margem}


def planilha_resposta(res):
    """Gera o Excel de resposta (item, quantidade, preço unitário, total) para conferir e importar no Ariba."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resposta"
    ws.append(["Evento", res["cabecalho"].get("evento") or ""])
    ws.append([])
    ws.append(["Item", "Descrição", "Quantidade", "Unidade", "Natureza", "Custo unitário", "Preço unitário", "Total", "Tributos no preço (%)"])
    for i, l in enumerate(res["itens"], 1):
        ws.append([i, l["descricao"], l["quantidade"], l["unidade"], l["natureza"], l["custo_unit"], l["preco_unit"], l["valor"], l["carga"]])
    ws.append([])
    ws.append(["", "Total da proposta", "", "", "", "", "", res["total"]])
    for k, v in res["retencoes"].items():
        ws.append(["", f"Retenção: {k}", "", "", "", "", "", -v])
    ws.append(["", "Líquido a receber", "", "", "", "", "", res["liquido"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
