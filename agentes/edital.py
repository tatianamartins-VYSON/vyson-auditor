"""Agente de edital: lê o edital (PDF, Word ou texto), analisa as exigências e gera a documentação modelo da proposta.

Passos:
  1. Leitor: extrai número, órgão, objeto, modalidade, data da sessão, valor estimado e critério de julgamento.
  2. Analista de habilitação: encontra as exigências (atestados, índices contábeis, capital mínimo, garantia, visita técnica,
     registro em conselho, certidões) e indica a página/trecho onde aparecem.
  3. Analista de riscos: aponta prazos curtos, exclusividade ME/EPP, garantia, amostra, consórcio vedado e subcontratação.
  4. Redator: gera em Word as declarações da Lei 14.133/2021 e a proposta comercial modelo, já com os dados da empresa.
"""
import io
import re
from datetime import date

EXIG = [
    ("Atestado de capacidade técnica", r"atestado[s]? de capacidade t[ée]cnica|capacidade t[ée]cnico-operacional|capacidade t[ée]cnico-profissional"),
    ("Registro no conselho (CREA, CAU, CRC…)", r"\b(CREA|CAU|CRC|CRA|CRM|OAB)\b"),
    ("Índices contábeis (LG, SG, LC)", r"liquidez geral|solv[êe]ncia geral|liquidez corrente|\bLG\b.{0,20}\bSG\b"),
    ("Capital social ou patrimônio líquido mínimo", r"(capital social|patrim[ôo]nio l[íi]quido) m[íi]nimo|n[ãa]o inferior a \d{1,2} ?%"),
    ("Balanço patrimonial", r"balan[çc]o patrimonial"),
    ("Garantia da proposta ou do contrato", r"garantia (da proposta|contratual|de execu[çc][ãa]o)|cau[çc][ãa]o|seguro-garantia|fian[çc]a banc[áa]ria"),
    ("Visita técnica", r"visita t[ée]cnica|vistoria"),
    ("Amostra ou prova de conceito", r"\bamostra|prova de conceito"),
    ("Certidões fiscais e trabalhistas", r"regularidade fiscal|CNDT|certid[ãa]o negativa|FGTS"),
    ("Planilha de custos / BDI", r"planilha (de custos|or[çc]ament[áa]ria)|\bBDI\b|composi[çc][ãa]o de custos"),
    ("Certificação (ISO, etc.)", r"\bISO ?\d{4,5}|certifica[çc][ãa]o"),
]
RISCOS = [
    ("Exclusiva para ME/EPP", r"exclusiv[ao] (para|a) (microempresas?|ME)|participa[çc][ãa]o exclusiva", "Só microempresas e empresas de pequeno porte podem participar (LC 123/2006, art. 48)."),
    ("Consórcio vedado", r"vedad[ao] a participa[çc][ãa]o de (empresas reunidas em )?cons[óo]rcio|n[ãa]o ser[áa] admitida a participa[çc][ãa]o de cons[óo]rcio", "Não é possível participar em consórcio."),
    ("Subcontratação limitada", r"subcontrata[çc][ãa]o", "O edital trata de subcontratação: confira o limite antes de contar com terceiros."),
    ("Margem de preferência", r"margem de prefer[êe]ncia", "Há margem de preferência para produtos nacionais (Lei 14.133/2021, art. 26)."),
    ("Reajuste e repactuação", r"reajuste|repactua[çc][ãa]o", "Confira o índice e a data-base do reajuste."),
]


def ler_texto(nome, dados):
    ext = nome.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(dados)) as pdf:
            return [(i + 1, p.extract_text() or "") for i, p in enumerate(pdf.pages[:300])]
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(dados))
        return [(1, "\n".join(p.text for p in d.paragraphs))]
    return [(1, dados.decode("utf-8", "ignore"))]


def analisar(nome, dados):
    paginas = ler_texto(nome, dados)
    tudo = "\n".join(t for _, t in paginas)
    g = lambda rx: (lambda m: m.group(1).strip() if m else None)(re.search(rx, tudo, re.I))
    cab = {
        "numero": g(r"(?:PREG[ÃA]O|CONCORR[ÊE]NCIA|EDITAL|DISPENSA)[^\n]{0,30}?N[º°o.]*\s*([\d./-]+)"),
        "modalidade": g(r"\b(Preg[ãa]o Eletr[ôo]nico|Concorr[êe]ncia(?: Eletr[ôo]nica)?|Dispensa Eletr[ôo]nica|Leil[ãa]o|Di[áa]logo Competitivo)\b"),
        "orgao": g(r"\n?([A-ZÀ-Ú][A-ZÀ-Ú ]{8,}(?:PREFEITURA|SECRETARIA|TRIBUNAL|MINIST[ÉE]RIO|UNIVERSIDADE|C[ÂA]MARA|GOVERNO)[A-ZÀ-Ú ]*|(?:PREFEITURA|SECRETARIA|TRIBUNAL|MINIST[ÉE]RIO|UNIVERSIDADE|C[ÂA]MARA|GOVERNO)[A-ZÀ-Ú ,]{4,80})"),
        "objeto": g(r"OBJETO\s*[:\-–]?\s*(?:\d[\d.]*\s*)?(.{20,400}?)(?:\.\s|\n\d)"),
        "sessao": g(r"(?:sess[ãa]o p[úu]blica|abertura das propostas|data da sess[ãa]o)[^\n]{0,60}?(\d{2}/\d{2}/\d{4}(?:[^\n]{0,15}?\d{1,2}[h:]\d{2})?)"),
        "valor": g(r"valor (?:total )?(?:estimado|global|m[áa]ximo)[^\n]{0,40}?R\$\s*(\d[\d.]*,\d{2})"),
        "criterio": g(r"(menor pre[çc]o|maior desconto|t[ée]cnica e pre[çc]o|melhor t[ée]cnica)"),
    }
    exig = []
    for nomex, rx in EXIG:
        pags = [p for p, t in paginas if re.search(rx, t, re.I)]
        if pags:
            m = re.search(r"[^.\n]{0,120}(" + rx + r")[^.\n]{0,160}", tudo, re.I)
            exig.append({"exigencia": nomex, "paginas": pags[:6], "trecho": re.sub(r"\s+", " ", m.group(0)).strip()[:260] if m else ""})
    riscos = [{"risco": n, "nota": nota} for n, rx, nota in RISCOS if re.search(rx, tudo, re.I)]
    UN = r"(UN|UND|UNID|UNIDADE|SC|SACO|KG|M2|M²|M3|M³|M|ML|L|LT|CX|CAIXA|PCT|PACOTE|PC|PÇ|PEÇA|DIA|DIÁRIA|MÊS|MES|SERV|SERVIÇO|SV|HORA|H|KIT|RESMA|GL|GALÃO)"
    itens = []
    for ln in tudo.splitlines():
        m = re.match(r"^\s*(\d{1,3})[\s.|)\-–]+(.{4,160}?)\s*[|;]?\s*" + UN + r"\.?\s*[|;]?\s*(\d+(?:[.,]\d+)?)\s*$", ln, re.I) or \
            re.match(r"^\s*(\d{1,3})[\s.|)\-–]+(.{4,160}?)\s*[|;]?\s*(\d+(?:[.,]\d+)?)\s*[|;]?\s*" + UN + r"\.?\s*$", ln, re.I)
        if m:
            a, b = (m.group(3), m.group(4)) if not m.group(3)[0].isdigit() else (m.group(4), m.group(3))
            try:
                itens.append({"item": int(m.group(1)), "descricao": m.group(2).strip(" |;-–"), "unidade": a.lower(), "quantidade": float(b.replace(".", "").replace(",", "."))})
            except ValueError:
                pass
    lg = re.search(r"liquidez geral[^\n]{0,40}?(\d[,.]\d{1,2})", tudo, re.I)
    if lg:
        riscos.append({"risco": "Índice mínimo", "nota": f"Liquidez geral mínima de {lg.group(1)}: confira no seu balanço antes de participar."})
    passos = [f"Leitor: {len(paginas)} página(s); edital {cab['numero'] or 'sem número'}, {cab['modalidade'] or 'modalidade não identificada'}",
              f"Analista de habilitação: {len(exig)} exigência(s) encontradas", f"Analista de riscos: {len(riscos)} ponto(s) de atenção",
              "Fabricante: mais de 30 documentos em papel timbrado (declarações, procuração, carta de apresentação, proposta comercial, cronograma e checklist) em Word"]
    passos.insert(1, f"Leitor de itens: {len(itens)} item(ns) do termo de referência")
    return {"cabecalho": cab, "exigencias": exig, "riscos": riscos, "itens": itens[:200], "passos": passos, "paginas": len(paginas)}


DECL = [
    ("Declaração de cumprimento dos requisitos de habilitação", "declara, sob as penas da lei, que cumpre plenamente os requisitos de habilitação exigidos no edital, conforme o art. 63, I, da Lei nº 14.133/2021."),
    ("Declaração de inexistência de fato impeditivo", "declara, sob as penas da lei, que não há fato impeditivo à sua habilitação e que não está impedida ou declarada inidônea para licitar ou contratar com a Administração Pública (Lei nº 14.133/2021, arts. 14 e 156)."),
    ("Declaração de não emprego de menor", "declara, para fins do art. 7º, XXXIII, da Constituição Federal e do art. 68, VI, da Lei nº 14.133/2021, que não emprega menor de 18 anos em trabalho noturno, perigoso ou insalubre, nem menor de 16 anos, salvo na condição de aprendiz, a partir de 14 anos."),
    ("Declaração de reserva de cargos para pessoa com deficiência", "declara que cumpre as exigências de reserva de cargos para pessoa com deficiência e para reabilitado da Previdência Social previstas em lei (Lei nº 14.133/2021, art. 63, IV; Lei nº 8.213/1991, art. 93)."),
    ("Declaração de elaboração independente da proposta", "declara que a proposta apresentada foi elaborada de maneira independente, sem consulta, acordo ou troca de informações com qualquer outro participante, e que seu conteúdo não foi informado a terceiros antes da abertura oficial."),
    ("Declaração de que a proposta compreende todos os custos", "declara que a proposta econômica compreende a integralidade dos custos para atendimento dos direitos trabalhistas assegurados na Constituição Federal, nas leis trabalhistas, nas normas infralegais, nas convenções coletivas de trabalho e nos termos de ajustamento de conduta vigentes na data de entrega da proposta (Lei nº 14.133/2021, art. 63, §1º)."),
]
DECL_EXTRA = [
    ("Declaração de conhecimento do edital e das condições locais", "declara que tomou conhecimento de todas as informações, condições locais e exigências do edital e de seus anexos, aceitando-as integralmente, e que não alegará desconhecimento para eximir-se de obrigações contratuais (Lei nº 14.133/2021, art. 63, §2º)."),
    ("Declaração de renúncia à visita técnica", "declara que optou por não realizar a vistoria prévia, assumindo integral responsabilidade pelas condições do local e pelo cumprimento das obrigações, sem direito a alegar desconhecimento das condições do objeto (Lei nº 14.133/2021, art. 63, §§2º e 3º)."),
    ("Declaração de inexistência de parentesco com agentes públicos", "declara que não possui, em seu quadro societário ou como representante, cônjuge, companheiro ou parente em linha reta, colateral ou por afinidade, até o terceiro grau, de agente público que atue na licitação ou na fiscalização e gestão do contrato (Lei nº 14.133/2021, art. 14, IV)."),
    ("Declaração de disponibilidade de pessoal, instalações e equipamentos", "declara que disporá, na data da contratação, das instalações, do aparelhamento e do pessoal técnico adequados e disponíveis para a execução do objeto (Lei nº 14.133/2021, art. 67, §8º)."),
    ("Declaração de responsabilidade socioambiental", "declara que observa a legislação ambiental e as normas de sustentabilidade aplicáveis ao objeto, incluindo a destinação adequada de resíduos, e que não utiliza trabalho degradante ou forçado (Lei nº 14.133/2021, arts. 5º e 11, IV; Constituição Federal, art. 225)."),
    ("Indicação de preposto", "indica como preposto, para representá-la na execução do contrato, o(a) Sr.(a) [NOME DO PREPOSTO], CPF nº [CPF], telefone [TELEFONE], e-mail [E-MAIL], com poderes para tratar de todas as questões relativas ao contrato (Lei nº 14.133/2021, art. 118)."),
    ("Carta de credenciamento", "credencia o(a) Sr.(a) [NOME DO CREDENCIADO], CPF nº [CPF], para representá-la na sessão pública da licitação, com poderes para formular lances, negociar preços, interpor e desistir de recursos e praticar todos os demais atos do certame."),
    ("Declaração de enquadramento tributário", "declara que é optante pelo regime tributário [SIMPLES NACIONAL / LUCRO PRESUMIDO / LUCRO REAL], ciente de que a proposta contempla os tributos desse regime e de que as retenções na fonte seguirão a legislação aplicável (IN RFB nº 1.234/2012; LC nº 123/2006, art. 18)."),
]
DECL_MAIS = [
    ("Declaração de vistoria técnica realizada", "declara que realizou a vistoria técnica no local de execução do objeto, por meio de seu responsável técnico, e que tomou conhecimento das condições e peculiaridades do local (Lei nº 14.133/2021, art. 63, §§2º e 3º)."),
    ("Procuração para representação na licitação", "nomeia e constitui seu bastante procurador o(a) Sr.(a) [NOME DO PROCURADOR], CPF nº [CPF], a quem confere poderes para representá-la na licitação, formular lances, negociar preços, interpor e desistir de recursos, assinar atas e praticar todos os demais atos do certame (Código Civil, arts. 653 e seguintes; Lei nº 14.133/2021, art. 63)."),
    ("Carta de apresentação da proposta", "apresenta sua proposta para o objeto da licitação em referência, declarando que examinou o edital e seus anexos, que concorda com todas as condições estabelecidas e que se compromete a executar o objeto nos prazos e condições exigidos (Lei nº 14.133/2021, art. 59)."),
    ("Declaração de não estar impedida, suspensa ou declarada inidônea", "declara que não consta nos cadastros de empresas inidôneas e suspensas (CEIS), de empresas punidas (CNEP) nem em qualquer registro de sanção que a impeça de licitar ou contratar com a Administração Pública (Lei nº 14.133/2021, arts. 14 e 156; Lei nº 12.846/2013, art. 22)."),
    ("Declaração de compromisso de manter as condições de habilitação", "compromete-se a manter, durante toda a execução do contrato, todas as condições de habilitação e qualificação exigidas na licitação (Lei nº 14.133/2021, art. 92, XVI)."),
    ("Declaração de compromisso de prestar garantia contratual", "compromete-se a prestar a garantia contratual exigida no edital, em uma das modalidades admitidas em lei, no prazo fixado para a assinatura do contrato (Lei nº 14.133/2021, arts. 96 a 102)."),
    ("Declaração de validade da proposta", "declara que sua proposta tem validade de 60 (sessenta) dias, contados da data da sessão pública, ou o prazo maior exigido no edital (Lei nº 14.133/2021, art. 59)."),
    ("Declaração de exequibilidade da proposta", "declara que os preços ofertados são exequíveis, cobrem todos os custos de execução e permitem o cumprimento integral do contrato, comprometendo-se a demonstrar a exequibilidade quando solicitado (Lei nº 14.133/2021, art. 59, III e §§2º a 4º)."),
    ("Declaração de inexistência de servidor do órgão no quadro da empresa", "declara que não possui em seu quadro societário, de dirigentes ou de empregados servidor ou dirigente do órgão ou entidade contratante ou responsável pela licitação (Lei nº 14.133/2021, art. 9º, §1º)."),
    ("Termo de confidencialidade e proteção de dados", "compromete-se a manter sigilo sobre as informações a que tiver acesso e a tratar dados pessoais somente para a execução do objeto, observando a Lei Geral de Proteção de Dados (Lei nº 13.709/2018)."),
    ("Declaração de cumprimento da cota de aprendizes", "declara que cumpre a cota de contratação de aprendizes prevista em lei, quando aplicável ao seu porte e atividade (CLT, art. 429; Lei nº 14.133/2021, art. 63, IV)."),
    ("Relação da equipe técnica e indicação do responsável técnico", "indica como responsável técnico pela execução do objeto o(a) profissional [NOME], registro profissional nº [REGISTRO], e apresenta a relação da equipe técnica que participará dos serviços (Lei nº 14.133/2021, art. 67, I e §6º)."),
    ("Declaração de programa de integridade", "declara que possui ou se compromete a implantar programa de integridade, quando exigido pelo edital em razão do valor da contratação (Lei nº 14.133/2021, art. 25, §4º; Decreto nº 12.304/2024)."),
    ("Declaração de marca, modelo e procedência dos bens", "declara a marca, o modelo, o fabricante e a procedência dos bens ofertados, conforme indicado na proposta, comprometendo-se a entregar exatamente os bens cotados (Lei nº 14.133/2021, art. 41)."),
]
DECL_ME = ("Declaração de enquadramento como ME/EPP", "declara, sob as penas da lei, que se enquadra como microempresa ou empresa de pequeno porte, nos termos do art. 3º da Lei Complementar nº 123/2006, e que não se encontra em nenhuma das vedações do §4º do mesmo artigo, estando apta a usufruir do tratamento favorecido previsto nos arts. 42 a 49.")


def gerar_docx(analise, empresa):
    """Gera um Word com todas as declarações e a proposta comercial modelo, preenchidas com os dados informados."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(11)
    cab = analise["cabecalho"]
    e = {k: (empresa.get(k) or f"[{k.upper()}]") for k in ("razao", "cnpj", "endereco", "representante", "cargo", "cidade")}
    ident = f"{e['razao']}, inscrita no CNPJ sob o nº {e['cnpj']}, com sede em {e['endereco']}, por seu representante legal {e['representante']}, {e['cargo']},"
    ref = f"Ref.: {cab.get('modalidade') or 'Licitação'} nº {cab.get('numero') or '____'}" + (f" — {cab['orgao']}" if cab.get("orgao") else "")
    hoje = date.today().strftime("%d/%m/%Y")
    # papel timbrado: cabeçalho e rodapé de todas as páginas com os dados da empresa (e o logotipo, se enviado)
    sec = d.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if empresa.get("logo_b64"):
        import base64
        try:
            hp.add_run().add_picture(io.BytesIO(base64.b64decode(empresa["logo_b64"])), height=Pt(36))
            hp.add_run("\n")
        except Exception:  # noqa: BLE001
            pass
    rh = hp.add_run(e["razao"])
    rh.bold = True
    hp.add_run(f"\nCNPJ {e['cnpj']} · {e['endereco']}").font.size = Pt(9)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run(f"{e['razao']} · {ref}").font.size = Pt(8)
    docs = DECL + ([DECL_ME] if empresa.get("me") else []) + DECL_EXTRA + DECL_MAIS
    if empresa.get("regime"):
        docs = [(t, c.replace("[SIMPLES NACIONAL / LUCRO PRESUMIDO / LUCRO REAL]", empresa["regime"].upper())) for t, c in docs]
    for i, (titulo, corpo) in enumerate(docs):
        if i:
            d.add_page_break()
        h = d.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run(titulo.upper())
        r.bold = True
        d.add_paragraph(ref)
        p = d.add_paragraph(f"{ident} {corpo}")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        d.add_paragraph(f"\n{e['cidade']}, {hoje}.\n\n\n______________________________\n{e['representante']}\n{e['cargo']}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()
    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.add_run("PROPOSTA COMERCIAL").bold = True
    d.add_paragraph(ref)
    d.add_paragraph(f"Objeto: {cab.get('objeto') or '[descrever o objeto conforme o edital]'}").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, ["Item", "Descrição", "Quantidade", "Valor unitário (R$)", "Valor total (R$)"]):
        c.text = txt
    for i in range(1, 4):
        for c, txt in zip(t.add_row().cells, [str(i), "", "", "", ""]):
            c.text = txt
    for linha in ["Valor global da proposta: R$ ____________ (por extenso).",
                  "Validade da proposta: 60 (sessenta) dias a contar da data da sessão pública.",
                  "Nos preços estão incluídos todos os custos diretos e indiretos, tributos, encargos sociais e trabalhistas, seguros, fretes e o BDI.",
                  f"Dados bancários: Banco ____ Agência ____ Conta ____, em nome de {e['razao']}."]:
        d.add_paragraph(linha).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_paragraph(f"\n{e['cidade']}, {hoje}.\n\n\n______________________________\n{e['representante']}\n{e['cargo']}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cronograma(d, analise, ref)
    _checklist(d, analise, ref, empresa)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


PASTAS = [
    ("Habilitação jurídica (art. 66)", ["Contrato social ou ato constitutivo e alterações", "Documento dos sócios e do representante", "Procuração ou carta de credenciamento"]),
    ("Regularidade fiscal, social e trabalhista (art. 68)", ["Cartão CNPJ", "CND federal (Receita/PGFN)", "Certidão estadual", "Certidão municipal", "Certificado do FGTS (CRF)", "CNDT (débitos trabalhistas)"]),
    ("Qualificação técnica (art. 67)", ["Atestados de capacidade técnica", "Registro no conselho profissional (CREA, CRC, CRA...)", "Declaração de disponibilidade de pessoal e equipamentos"]),
    ("Qualificação econômico-financeira (art. 69)", ["Balanço patrimonial e DRE dos 2 últimos exercícios", "Certidão negativa de falência", "Índices de liquidez (LG, SG, LC)"]),
]


def _titulo(d, txt, ref):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d.add_page_break()
    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.add_run(txt).bold = True
    d.add_paragraph(ref)


def _tabela(d, cab, linhas):
    t = d.add_table(rows=1, cols=len(cab))
    t.style = "Table Grid"
    for c, x in zip(t.rows[0].cells, cab):
        c.text = x
    for l in linhas:
        for c, x in zip(t.add_row().cells, l):
            c.text = str(x)


def _cronograma(d, analise, ref):
    _titulo(d, "CRONOGRAMA FÍSICO-FINANCEIRO", ref)
    itens = [i.get("descricao", "") for i in analise.get("itens", [])][:30] or ["Mobilização e serviços preliminares", "Execução principal", "Acabamento, testes e entrega"]
    _tabela(d, ["Etapa / item", "Valor (R$)", "Mês 1 (%)", "Mês 2 (%)", "Mês 3 (%)", "Total (%)"], [[x, "", "", "", "", "100"] for x in itens] + [["TOTAL", "", "", "", "", "100"]])
    d.add_paragraph("Ajuste os meses ao prazo de execução do edital. A soma dos percentuais de cada linha deve ser 100%.")


def _checklist(d, analise, ref, empresa):
    _titulo(d, "CHECKLIST DE HABILITAÇÃO E ÍNDICE DA DOCUMENTAÇÃO", ref)
    linhas = []
    for pasta, docs in PASTAS:
        linhas += [[pasta, x, "", ""] for x in docs]
    linhas += [["Declarações", t, "", ""] for t, _ in DECL + ([DECL_ME] if empresa.get("me") else []) + DECL_EXTRA]
    linhas += [["Edital (pág. " + ", ".join(map(str, x["paginas"])) + ")", x["exigencia"], "", ""] for x in analise.get("exigencias", [])]
    _tabela(d, ["Pasta", "Documento", "Situação (ok / falta / vencido)", "Validade"], linhas)
    d.add_paragraph("Lei nº 14.133/2021, arts. 62 a 70. Confira cada certidão no site do órgão emissor no dia do envio.")


# ---------------- validação dos documentos preenchidos contra o edital ----------------
TIPOS_DOC = [
    ("Declaração de cumprimento dos requisitos de habilitação", r"cumpre plenamente os requisitos de habilita"),
    ("Declaração de inexistência de fato impeditivo", r"fato impeditivo"),
    ("Declaração de não emprego de menor", r"menor de (18|dezoito)"),
    ("Declaração de reserva de cargos para pessoa com deficiência", r"pessoa com defici[êe]ncia"),
    ("Declaração de elaboração independente da proposta", r"elaborad[ao] de maneira independente"),
    ("Declaração de que a proposta compreende todos os custos", r"integralidade dos custos"),
    ("Declaração de enquadramento como ME/EPP", r"microempresa ou empresa de pequeno porte"),
    ("Declaração de conhecimento do edital", r"tomou conhecimento de todas as informa"),
    ("Declaração de inexistência de parentesco", r"parente em linha reta"),
    ("Indicação de preposto", r"indica como preposto"),
    ("Carta de credenciamento", r"\bcredencia o"),
    ("Proposta comercial", r"proposta comercial|valor global da proposta"),
    ("CND federal (Receita/PGFN)", r"tributos federais e [àa] d[íi]vida ativa da uni[ãa]o"),
    ("Certificado do FGTS (CRF)", r"certificado de regularidade do fgts|\bCRF\b"),
    ("CNDT (débitos trabalhistas)", r"certid[ãa]o negativa de d[ée]bitos trabalhistas"),
    ("Balanço patrimonial", r"balan[çc]o patrimonial"),
    ("Atestado de capacidade técnica", r"atesta(mos)? (para os devidos fins )?que|atestado de capacidade"),
    ("Contrato social", r"contrato social|ato constitutivo"),
]


# fonte legal e onde emitir cada documento (dinheiro do cliente: tudo com base e link oficial)
FONTE_DOC = {
    "CND federal (Receita/PGFN)": ("Lei 14.133/2021, art. 68, III; Portaria Conjunta RFB/PGFN 1.751/2014", "https://solucoes.receita.fazenda.gov.br/Servicos/certidaointernet/PJ/Emitir"),
    "Certificado do FGTS (CRF)": ("Lei 14.133/2021, art. 68, IV; Lei 8.036/1990, art. 27", "https://consulta-crf.caixa.gov.br/consultacrf/pages/consultaEmpregador.jsf"),
    "CNDT (débitos trabalhistas)": ("Lei 14.133/2021, art. 68, V; CLT, art. 642-A", "https://cndt-certidao.tst.jus.br/inicio.faces"),
    "Balanço patrimonial": ("Lei 14.133/2021, art. 69, I", ""),
    "Atestado de capacidade técnica": ("Lei 14.133/2021, art. 67, II", ""),
    "Contrato social": ("Lei 14.133/2021, art. 66", "https://www.gov.br/empresas-e-negocios/pt-br/redesim"),
    "Proposta comercial": ("Lei 14.133/2021, arts. 59 e 63", ""),
    "Declaração de enquadramento como ME/EPP": ("LC 123/2006, arts. 3º e 42 a 49; Lei 14.133/2021, art. 4º", ""),
}
FONTE_DECL = "Lei 14.133/2021, art. 63 (declarações de habilitação)"
SICAF = "https://www.gov.br/compras/pt-br/sistemas/conheca-o-compras/sicaf-digital"


def _data_sessao(analise):
    m = re.search(r"(\d{2})/(\d{2})/(\d{4})", (analise.get("cabecalho") or {}).get("sessao") or "")
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None
    except ValueError:
        return None


def validar(analise, docs, empresa):
    """docs: lista de (nome, bytes). Confere presença, CNPJ, referência ao edital, validade e campos em branco."""
    num = re.sub(r"\D", "", (analise["cabecalho"].get("numero") or ""))
    cnpj = re.sub(r"\D", "", empresa.get("cnpj") or "")
    hoje = date.today()
    sessao = _data_sessao(analise)
    achados, presentes = [], {}
    for nome, dados in docs:
        texto = "\n".join(t for _, t in ler_texto(nome, dados))
        tipos = [t for t, rx in TIPOS_DOC if re.search(rx, texto, re.I)]
        tipo = tipos[0] if tipos else None
        probs = []
        if not tipo:
            probs.append("Tipo de documento não reconhecido")
        cn = [re.sub(r"\D", "", c) for c in re.findall(r"\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}", texto)]
        if cnpj and cn and cnpj not in cn:
            probs.append("CNPJ do documento diferente do CNPJ da empresa")
        if tipo and tipo.startswith(("Declaração", "Proposta")):
            if num and num not in re.sub(r"\D", "", texto):
                probs.append(f"Não cita o número do edital ({analise['cabecalho'].get('numero')})")
            if re.search(r"\[[A-ZÇÃ_ ]{3,}\]|_{6,}", texto):
                probs.append("Há campos em branco para preencher")
        val = re.search(r"(?:v[áa]lid[ao] at[ée]|validade[: ]+(?:at[ée] )?)\s*(\d{2})/(\d{2})/(\d{4})", texto, re.I)
        if val:
            try:
                dv = date(int(val.group(3)), int(val.group(2)), int(val.group(1)))
                if dv < hoje:
                    probs.append(f"Vencido em {dv.strftime('%d/%m/%Y')}")
                elif sessao and dv < sessao:
                    probs.append(f"Válido até {dv.strftime('%d/%m/%Y')}, antes da sessão de {sessao.strftime('%d/%m/%Y')}: emita outro antes da sessão")
                elif (dv - hoje).days <= 10:
                    probs.append(f"Vence em {dv.strftime('%d/%m/%Y')}: renove antes da sessão")
            except ValueError:
                pass
        for t in tipos:
            presentes[t] = presentes.get(t, 0) + 1
        fonte, link = FONTE_DOC.get(tipo, (FONTE_DECL if tipo and tipo.startswith(("Declaração", "Indicação", "Carta")) else "", "")) if tipo else ("", "")
        achados.append({"arquivo": nome, "tipo": ", ".join(tipos) if len(tipos) <= 2 else f"{len(tipos)} documentos no mesmo arquivo", "status": "ok" if not probs else "corrigir",
                        "problemas": probs, "fonte": fonte, "onde_emitir": link})
    exigidos = [t for t, _ in DECL] + ["Proposta comercial", "CND federal (Receita/PGFN)", "Certificado do FGTS (CRF)", "CNDT (débitos trabalhistas)"]
    if empresa.get("me"):
        exigidos.append(DECL_ME[0])
    for x in analise.get("exigencias", []):
        if x["exigencia"] in ("Balanço patrimonial", "Atestado de capacidade técnica"):
            exigidos.append(x["exigencia"])
    faltando = [t for t in dict.fromkeys(exigidos) if t not in presentes]
    ok = sum(a["status"] == "ok" for a in achados)
    falt = [{"documento": t, "fonte": FONTE_DOC.get(t, (FONTE_DECL, ""))[0], "onde_emitir": FONTE_DOC.get(t, ("", ""))[1]} for t in faltando]
    pronto = not faltando and ok == len(achados)
    return {"documentos": achados, "faltando": faltando, "faltando_detalhe": falt, "pronto_para_participar": pronto,
            "veredito": "Pronto para participar: todos os documentos exigidos estão presentes e válidos." if pronto else
                        f"NÃO está pronto: {len(achados) - ok} documento(s) para corrigir e {len(faltando)} faltando. Um documento ausente ou vencido inabilita a proposta (Lei 14.133/2021, art. 62 a 70).",
                        "sicaf": SICAF, "sessao": sessao.strftime("%d/%m/%Y") if sessao else None,
            "resumo": f"{ok} de {len(achados)} documento(s) sem pendência; {len(faltando)} exigido(s) ainda não enviado(s)"}
