"""Extrai os campos da nota fiscal de serviço a partir de XML (NFS-e Nacional ou ABRASF), PDF ou DOCX."""
import io
import re
import xml.etree.ElementTree as ET

CAMPOS = ("numero", "emissao", "prestador_cnpj", "prestador_nome", "tomador_cnpj", "tomador_nome", "municipio",
          "item_lc116", "ctribnac", "cod_municipal", "descricao", "valor_servicos", "deducoes", "aliquota_iss",
          "valor_iss", "iss_retido", "inss_retido", "valor_liquido")


def _num(s):
    if s is None or s == "":
        return None
    s = str(s).strip().replace("R$", "").replace("%", "").strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _xml(texto: str) -> dict:
    raiz = ET.fromstring(texto)
    todos = {}
    for el in raiz.iter():
        nome = el.tag.split("}")[-1]
        if el.text and el.text.strip() and nome not in todos:
            todos[nome] = el.text.strip()
    g = lambda *ns: next((todos[n] for n in ns if n in todos), None)
    ctrib = g("cTribNac")
    item = g("ItemListaServico")
    if not item and ctrib:
        item = f"{int(ctrib[:2])}.{ctrib[2:4]}"
    aliq = _num(g("pAliq", "Aliquota"))
    if aliq is not None and aliq < 1:  # ABRASF usa 0.05 para 5%
        aliq *= 100
    if "tpRetISSQN" in todos:  # NFS-e Nacional: 1 = não retido, 2 = tomador, 3 = intermediário
        retido = todos["tpRetISSQN"] in ("2", "3")
    elif "IssRetido" in todos:  # ABRASF: 1 = sim, 2 = não
        retido = todos["IssRetido"] == "1"
    else:
        retido = None
    tom = None
    for el in raiz.iter():
        t = el.tag.split("}")[-1].lower()
        if t.startswith("toma") or t == "tomadorservico":
            c = next((x.text.strip() for x in el.iter() if x.tag.split("}")[-1] in ("Cnpj", "CNPJ", "Cpf", "CPF") and x.text), None)
            if c:
                tom = c
                break
    op = g("OptanteSimplesNacional", "opSimpNac")
    simples = (op == "1") if "OptanteSimplesNacional" in todos else (op in ("2", "3") if op else None)
    ref = {}
    for el in raiz.iter():
        t = el.tag.split("}")[-1]
        if t in ("IBSCBS", "gIBSCBS"):
            for x in el.iter():
                n = x.tag.split("}")[-1]
                if n in ("CST", "cClassTrib", "vBC", "vIBS", "vCBS", "pIBSUF", "pIBSMun", "pCBS", "vIBSUF", "vIBSMun") and x.text and n not in ref:
                    ref[n] = x.text.strip()
    extra = {"atividade": g("DescricaoCodigoTributacaoMunicipio", "xTribMun", "xTribNac"), "ref_nbs": g("CodigoNbs", "cNBS"), "ref_cindop": g("cIndOp"), "ref_cst": ref.get("CST"), "ref_cclass": ref.get("cClassTrib"),
             "ref_vbc": _num(ref.get("vBC")), "ref_vibs": _num(ref.get("vIBS")), "ref_vcbs": _num(ref.get("vCBS")),
             "pis_valor": _num(g("vPis")), "cofins_valor": _num(g("vCofins")),
             "tomador_cnpj": tom, "simples": simples, "ir_retido": _num(g("ValorIr", "vRetIRRF")), "csll_retido": _num(g("ValorCsll", "vRetCSLL")),
             "pis_retido": _num(g("ValorPis")), "cofins_retido": _num(g("ValorCofins"))}
    return extra | {k: v for k, v in {
        "numero": g("nNFSe", "Numero", "nDPS"), "emissao": g("dhEmi", "DataEmissao", "dCompet"),
        "prestador_cnpj": g("CNPJ", "Cnpj"), "prestador_nome": g("xNome", "RazaoSocial"),
        "tomador_cnpj": None, "tomador_nome": None,
        "municipio": g("cLocPrestacao", "CodigoMunicipio", "cLocIncid"),
        "item_lc116": item, "ctribnac": ctrib, "cod_municipal": g("cTribMun", "CodigoTributacaoMunicipio"),
        "descricao": g("xDescServ", "Discriminacao"),
        "valor_servicos": _num(g("vServ", "ValorServicos")), "deducoes": _num(g("vDR", "ValorDeducoes")),
        "aliquota_iss": aliq, "valor_iss": _num(g("vISSQN", "ValorIss")),
        "iss_retido": retido,
        "inss_retido": _num(g("vRetCP", "ValorInss")), "valor_liquido": _num(g("vLiq", "ValorLiquidoNfse")),
    }.items() if not (k == "tomador_cnpj" and v is None)}


ROTULOS = {
    "numero": r"(?:N[úu]mero da (?:Nota|NFS-?e)|NFS-?e n[ºo°]?)\s*[:\-]?\s*(\d+)",
    "emissao": r"(?:Data (?:e Hora )?(?:de )?Emiss[ãa]o)\s*[:\-]?\s*([\d/]{10})",
    "item_lc116": r"Item[^\n:]*:\s*(\d{1,2}\.\d{2})\b",
    "cod_municipal": r"C[óo]digo (?:de )?(?:Tributa[çc][ãa]o|Servi[çc]o) Municipal\s*[:\-]?\s*([\d.\-/]+)",
    "valor_servicos": r"Valor (?:Total )?(?:dos )?Servi[çc]os\s*[:\-]?\s*R?\$?\s*([\d.,]+)",
    "deducoes": r"Dedu[çc][õo]es\s*[:\-]?\s*R?\$?\s*([\d.,]+)",
    "aliquota_iss": r"Al[íi]quota(?: do ISS)?\s*(?:\(%\))?\s*[:\-]?\s*([\d.,]+)\s*%?",
    "valor_iss": r"Valor (?:do )?ISS(?:QN)?\s*[:\-]?\s*R?\$?\s*([\d.,]+)",
    "inss_retido": r"INSS(?: Retido)?\s*[:\-]?\s*R?\$?\s*([\d.,]+)",
    "valor_liquido": r"Valor L[íi]quido(?: da Nota)?\s*[:\-]?\s*R?\$?\s*([\d.,]+)",
    "municipio": r"Munic[íi]pio(?: de Incid[êe]ncia| da Presta[çc][ãa]o)?\s*[:\-]?\s*([A-ZÀ-Úa-zà-ú ]+?(?:/[A-Z]{2})?)\s*(?:\n|$)",
    "descricao": r"(?:Discrimina[çc][ãa]o|Descri[çc][ãa]o) (?:dos |do )?Servi[çc]os?\s*[:\-]?\s*(.+)",
}


def _giss(texto: str) -> dict | None:
    """Layout da NFS-e do GISS (Maceió e outros municípios): 'Atividade Econômica 17.01 / 8211300 - ...'."""
    if not re.search(r"giss\.com\.br|Atividade Econ[ôo]mica", texto, re.I):
        return None
    f = lambda rx, n=1: (lambda m: m.group(n).strip() if m else None)(re.search(rx, texto, re.I))
    marca = lambda rx: (lambda m: "X" in m.group(1).upper() if m else None)(re.search(rx, texto, re.I))
    cn = re.findall(r"CPF/CNPJ:\s*([\d./-]{11,18})", texto)
    fed = re.search(r"PIS\s+Tipo PIS\s+COFINS.*?\n\s*([\d.,]+)\s+\S.*?\s([\d.,]+)\s+\S.*?\s([\d.,]+)\s+([\d.,]+)\s+([\d.,]+)", texto, re.I)
    mun = f(r"Mun[íi]c[íi]pio Incid[êe]ncia\s+([A-ZÀ-Ú ]+)\s*-\s*([A-Z]{2})")
    d = {k: None for k in CAMPOS}
    d.update({"numero": f(r"NFS-?e\s*\n\s*(\d+)\s*\n") or f(r"NFS-?e\s+(\d+)\b"), "emissao": f(r"Emiss[ãa]o da NFS-?e\s*\n?\s*([\d/]{10})"),
              "prestador_cnpj": cn[0] if cn else None, "tomador_cnpj": cn[1] if len(cn) > 1 else None,
              "item_lc116": (lambda v: v and f"{int(v.split('.')[0])}.{v.split('.')[1]}")(f(r"Atividade Econ[ôo]mica\s*\n?\s*(\d{1,2}\.\d{2})")),
              "cod_municipal": f(r"Atividade Econ[ôo]mica\s*\n?\s*\d{1,2}\.\d{2}\s*/\s*(\d+)"),
              "descricao": f(r"Discrimina[çc][ãa]o do Servi[çc]o\s*\n(.+)"),
              "valor_servicos": _num(f(r"Valor do Servi[çc]o\s+([\d.]+,\d{2})")), "deducoes": _num(f(r"Dedu[çc][õo]es Previstas em Lei\s+([\d.]+,\d{2})")),
              "aliquota_iss": _num(f(r"Al[íi]quota\s+([\d.]+,\d{2})")), "valor_iss": _num(f(r"\bISSQN\s+([\d.]+,\d{2})")),
              "valor_liquido": _num(f(r"Valor L[íi]quido\s+([\d.]+,\d{2})")),
              "municipio": (mun.title() if mun else None) and ("Maceió/AL" if "MACEI" in mun.upper() else mun),
              "simples": marca(r"Simples Nacional\s*\(\s*(X?)\s*\)\s*Sim"), "iss_retido": marca(r"ISSQN a Reter\s*\(\s*(X?)\s*\)\s*Sim"),
              "inss_retido": _num(fed.group(3)) if fed else None, "ir_retido": _num(fed.group(4)) if fed else None, "csll_retido": _num(fed.group(5)) if fed else None,
              "pis_retido": _num(fed.group(1)) if fed else None, "cofins_retido": _num(fed.group(2)) if fed else None,
              "isento": False, "locacao_bens": False,
              "atividade": f(r"Atividade Econ[ôo]mica\s*\n?\s*\d{1,2}\.\d{2}\s*/\s*\d+\s*-\s*(.+)"),
              "ref_nbs": f(r"Nomenclatura Brasileira de Servi[çc]os\s+([\d.]+)"), "ref_cindop": f(r"Indicador de Opera[çc][ãa]o\s+(\d{6})"),
              "ref_cst": f(r"Situa[çc][ãa]o Tribut[áa]ria\s+(\d{3})"), "ref_cclass": f(r"Classifica[çc][ãa]o Tribut[áa]ria\s+(\d{6})"),
              "ref_vibs": _num(f(r"\bIBS\s+([\d.]+,\d{2})")), "ref_vcbs": _num(f(r"\bCBS\s+([\d.]+,\d{2})"))})
    return d


def _texto(texto: str) -> dict:
    g = _giss(texto)
    if g:
        return g
    d = {k: None for k in CAMPOS}
    for k, rx in ROTULOS.items():
        m = re.search(rx, texto, re.I)
        if m:
            v = m.group(1).strip()
            d[k] = _num(v) if k in ("valor_servicos", "deducoes", "aliquota_iss", "valor_iss", "inss_retido", "valor_liquido") else v
    # tabela em duas linhas: rótulos numa linha e valores na seguinte (layout comum de prefeituras)
    linhas = texto.splitlines()
    TAB = [("valor_liquido", r"Valor L[íi]quido"), ("deducoes", r"Dedu[çc][õo]es"), ("base", r"Base de C[áa]lculo"),
           ("aliquota_iss", r"Al[íi]quota"), ("valor_iss", r"Valor do ISS"), ("valor_servicos", r"Valor (?:Total )?(?:dos )?Servi[çc]os")]
    for i, ln in enumerate(linhas[:-1]):
        achados = sorted((m.start(), k) for k, rx in TAB for m in [re.search(rx, ln, re.I)] if m)
        nums = re.findall(r"\d[\d.]*,\d{2}", linhas[i + 1])
        if len(achados) >= 3 and len(nums) == len(achados):
            for (_, k), v in zip(achados, nums):
                d[k] = _num(v)
    sec = lambda nome: re.search(nome + r"[\s\S]{0,300}?(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})", texto, re.I)
    p, t = sec(r"PRESTADOR"), sec(r"TOMADOR")
    if p:
        d["prestador_cnpj"] = p.group(1)
    if t:
        d["tomador_cnpj"] = t.group(1)
    if not d.get("numero"):
        m = re.search(r"N[úu]mero da Nota:?[\s\S]{0,120}?\b(\d{6,15})\b", texto, re.I)
        d["numero"] = m.group(1) if m else None
    if d.get("valor_servicos") is None:
        m = re.search(r"VALOR TOTAL DA NOTA\s*=?\s*R?\$?\s*([\d.,]+)", texto, re.I)
        d["valor_servicos"] = _num(m.group(1)) if m else d.get("base")
    d["simples"] = bool(re.search(r"Simples Nacional", texto, re.I))
    d["isento"] = bool(re.search(r"\b(ISEN[ÇC][ÃA]O|Isentos?/Imunes|sem a incid[êe]ncia de ISS)", texto, re.I))
    d["locacao_bens"] = bool(re.search(r"(LOCA[ÇC][ÃA]O|ALUGUEL) DE (M[ÁA]QUINAS|EQUIPAMENTOS|BENS)", texto, re.I))
    cnpjs = [] if p else re.findall(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", texto)
    if cnpjs:
        d["prestador_cnpj"] = cnpjs[0]
        d["tomador_cnpj"] = cnpjs[1] if len(cnpjs) > 1 else None
    d["iss_retido"] = bool(re.search(r"ISS\s*Retido\s*[:\-]?\s*(Sim|S\b)", texto, re.I)) if re.search(r"ISS\s*Retido", texto, re.I) else None
    return d


def extrair(nome: str, dados: bytes) -> dict:
    ext = nome.lower().rsplit(".", 1)[-1]
    if ext == "xml" or dados.lstrip()[:1] == b"<":
        d, origem = _xml(dados.decode("utf-8", "ignore")), "XML"
    elif ext == "pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(dados)) as pdf:
            texto = "\n".join(p.extract_text() or "" for p in pdf.pages)
        d, origem = _texto(texto), "PDF"
    elif ext == "docx":
        import docx
        doc = docx.Document(io.BytesIO(dados))
        linhas = [p.text for p in doc.paragraphs] + [" ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
        d, origem = _texto("\n".join(linhas)), "DOCX"
    else:
        d, origem = _texto(dados.decode("utf-8", "ignore")), "texto"
    d["origem"] = origem
    return d
